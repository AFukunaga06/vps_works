"""ファイル → テキスト変換（PDF / Word(.docx) / テキスト(.txt)）

拡張子で抽出方法を振り分ける:
  .pdf         → pdfplumber（テキスト埋め込み型のみ。画像PDFは空に近い）
  .docx        → python-docx（段落 + 表セルを抽出）
  .txt / .text → そのまま読み込み（UTF-8 / CP932 を順に試行）

旧バイナリ .doc は本モジュールでは未対応（変換ツールが必要）。

公開関数 extract_text() は従来通り。呼び出し側（main.py）は変更不要。
"""
from __future__ import annotations

import logging
from pathlib import Path

import pdfplumber

log = logging.getLogger("ai_server.pdf_parser")

# 受け付ける拡張子（小文字）
SUPPORTED_EXTS = {".pdf", ".docx", ".txt", ".text"}


def extract_text(file_path: str | Path) -> str:
    """拡張子に応じてテキストを抽出して返す。失敗時は例外を上げる。"""
    fp = Path(file_path)
    if not fp.is_file():
        raise FileNotFoundError(f"file not found: {fp}")

    ext = fp.suffix.lower()
    if ext == ".pdf":
        text = _extract_pdf(fp)
    elif ext == ".docx":
        text = _extract_docx(fp)
    elif ext in (".txt", ".text"):
        text = _extract_txt(fp)
    elif ext == ".doc":
        raise ValueError("旧形式の .doc は未対応です。.docx か PDF で保存し直してください。")
    else:
        raise ValueError(f"未対応のファイル形式です: {ext or '(拡張子なし)'}")

    text = text.strip()
    log.info("extracted %d chars from %s (%s)", len(text), fp.name, ext)
    return text


def _extract_pdf(fp: Path) -> str:
    """PDF 全ページのテキストを連結。"""
    pages: list[str] = []
    with pdfplumber.open(fp) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            try:
                t = page.extract_text() or ""
            except Exception as e:
                log.warning("page %d extract failed: %s", i, e)
                t = ""
            pages.append(t)
    return "\n".join(pages)


def _extract_docx(fp: Path) -> str:
    """Word(.docx) の段落と表セルを抽出。"""
    from docx import Document  # 遅延 import（PDF/TXT だけの時に不要）

    doc = Document(str(fp))
    parts: list[str] = [p.text for p in doc.paragraphs]
    # 表の中身も拾う（契約書では別表・料金表などが表で書かれることがある）
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            line = "\t".join(x for x in cells if x)
            if line:
                parts.append(line)
    return "\n".join(parts)


def _extract_txt(fp: Path) -> str:
    """テキストファイルを UTF-8 → CP932 の順で読み込む。"""
    data = fp.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "cp932"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    # 最後の手段: 不正バイトは置換して読む
    return data.decode("utf-8", errors="replace")
