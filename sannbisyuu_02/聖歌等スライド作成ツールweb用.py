"""
聖歌等スライド作成ツール Web版
Flask + python-pptx
"""
from flask import Flask, render_template, request, jsonify, send_file
import os, re, copy, io
from datetime import date
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches
from pptx.oxml.ns import qn

app = Flask(__name__)

BASE_DIR     = os.path.dirname(os.path.abspath(__file__))
DATA_DIR     = os.path.join(BASE_DIR, "data")
SANTABI_DIR  = os.path.join(DATA_DIR, "賛美集にない賛美")
OUTPUT_DIR   = os.path.join(BASE_DIR, "output")

os.makedirs(OUTPUT_DIR, exist_ok=True)
os.makedirs(DATA_DIR, exist_ok=True)
os.makedirs(SANTABI_DIR, exist_ok=True)

FULLWIDTH_DIGITS = str.maketrans("０１２３４５６７８９", "0123456789")

SEIKA_FILE_MAP = [
    ((1,   99),  "聖歌１～９９番.pptx"),
    ((100, 199), "聖歌１００番代.pptx"),
    ((200, 299), "聖歌２００番代.pptx"),
    ((300, 399), "聖歌３００番代.pptx"),
    ((400, 499), "聖歌４００番代.pptx"),
    ((500, 599), "聖歌５００番代.pptx"),
    ((600, 699), "聖歌６００番代.pptx"),
    ((700, 799), "聖歌７００番代.pptx"),
    ((800, 899), "聖歌８００番代.pptx"),
]

KYUSAN_FILE_MAP = [
    ((1,   100), "旧賛美集Ⅰ(１～100).pptx"),
    ((101, 158), "旧賛美集Ⅱ(101～158).pptx"),
]

NAVY  = RGBColor(0,   0, 128)
BLUE  = RGBColor(0,   0, 204)
BLACK = RGBColor(0,   0,   0)


def get_seika_file(n: int):
    for (lo, hi), fname in SEIKA_FILE_MAP:
        if lo <= n <= hi:
            return os.path.join(DATA_DIR, fname)
    return None


def get_kyusan_file(n: int):
    for (lo, hi), fname in KYUSAN_FILE_MAP:
        if lo <= n <= hi:
            return os.path.join(DATA_DIR, fname)
    return None


def get_santabi_items() -> list:
    result = []
    if not os.path.isdir(SANTABI_DIR):
        return result
    for fname in os.listdir(SANTABI_DIR):
        if os.path.splitext(fname)[1].lower() != ".pptx":
            continue
        base = os.path.splitext(fname)[0]
        prefix, title = base.split("_", 1) if "_" in base else ("", base)
        result.append({"prefix": prefix, "title": title})

    def sort_key(item):
        p = item["prefix"]
        # prefixが英字（ASCII）なら group=0（先頭）、日本語なら group=1
        is_english = bool(p) and all(ord(c) < 128 for c in p)
        group = 0 if is_english else 1
        return (group, p.lower() if is_english else p, item["title"])

    return sorted(result, key=sort_key)


def slide_text(slide) -> str:
    return "\n".join(
        sh.text_frame.text for sh in slide.shapes if sh.has_text_frame
    )


def extract_hymn_number(text: str):
    t = text.translate(FULLWIDTH_DIGITS)
    m = re.search(r"聖[\s　]*([\d\s　]{1,8})", t)
    if not m:
        return None
    digits = re.sub(r"[\s　]+", "", m.group(1))[:4]
    try:
        return int(digits)
    except ValueError:
        return None


def extract_kyusan_number(text: str):
    t = text.translate(FULLWIDTH_DIGITS)
    for pat in [r"賛[\s　]*([\d\s　]{1,4})",
                r"No[\s.]*(\d{1,4})",
                r"^\s*(\d{1,4})\s"]:
        m = re.search(pat, t, re.MULTILINE)
        if m:
            digits = re.sub(r"[\s　]+", "", m.group(1))[:4]
            try:
                return int(digits)
            except ValueError:
                pass
    return None


_NS_A = 'http://schemas.openxmlformats.org/drawingml/2006/main'

def _scale_spTree(sp_tree, sx: float, sy: float):
    """spTree内の全図形の位置・サイズを sx/sy 倍にスケーリングする。"""
    for xfrm in sp_tree.findall(f'.//{{{_NS_A}}}xfrm'):
        for tag, fx, fy in [('off','x','y'), ('chOff','x','y')]:
            el = xfrm.find(f'{{{_NS_A}}}{tag}')
            if el is not None:
                el.set('x', str(int(int(el.get('x', 0)) * sx)))
                el.set('y', str(int(int(el.get('y', 0)) * sy)))
        for tag, fx, fy in [('ext','cx','cy'), ('chExt','cx','cy')]:
            el = xfrm.find(f'{{{_NS_A}}}{tag}')
            if el is not None:
                el.set('cx', str(int(int(el.get('cx', 0)) * sx)))
                el.set('cy', str(int(int(el.get('cy', 0)) * sy)))
    # フォントサイズを縦方向の比率でスケーリング
    for tag in ('rPr', 'defRPr', 'endParaRPr'):
        for el in sp_tree.findall(f'.//{{{_NS_A}}}{tag}'):
            sz = el.get('sz')
            if sz:
                el.set('sz', str(int(int(sz) * sy)))


def copy_slide(src_slide, dst_prs, src_prs=None):
    blank  = dst_prs.slide_layouts[6]
    new_sl = dst_prs.slides.add_slide(blank)

    # 図形ツリーをコピー
    sp     = new_sl.shapes._spTree
    src_sp = src_slide.shapes._spTree
    sp.clear()
    for k, v in src_sp.attrib.items():
        sp.set(k, v)
    for ch in src_sp:
        sp.append(copy.deepcopy(ch))

    # ソースと出力のスライドサイズが異なる場合はスケーリング
    if src_prs is not None:
        src_w, src_h = src_prs.slide_width, src_prs.slide_height
        dst_w, dst_h = dst_prs.slide_width,  dst_prs.slide_height
        if src_w != dst_w or src_h != dst_h:
            _scale_spTree(sp, dst_w / src_w, dst_h / src_h)

    # 背景XML（<p:bg>）をコピー
    src_cSld = src_slide.element.find(qn('p:cSld'))
    dst_cSld = new_sl.element.find(qn('p:cSld'))
    if src_cSld is not None and dst_cSld is not None:
        src_bg = src_cSld.find(qn('p:bg'))
        if src_bg is not None:
            dst_bg = dst_cSld.find(qn('p:bg'))
            if dst_bg is not None:
                dst_cSld.remove(dst_bg)
            dst_spTree = dst_cSld.find(qn('p:spTree'))
            dst_cSld.insert(list(dst_cSld).index(dst_spTree), copy.deepcopy(src_bg))

    return new_sl


def set_bg(slide, color: RGBColor):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_black_slide(prs):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(sl, BLACK)
    sp = sl.shapes._spTree
    for ch in list(sp)[2:]:
        sp.remove(ch)
    return sl


def get_next_filename() -> str:
    today = date.today().strftime("%Y%m%d")
    for n in range(1, 1000):
        fname = f"{today}_{n:02d}.pptx"
        if not os.path.exists(os.path.join(OUTPUT_DIR, fname)):
            return fname
    return f"{today}_01.pptx"


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/santabi")
def api_santabi():
    return jsonify(get_santabi_items())


@app.route("/api/next_filename")
def api_next_filename():
    return jsonify({"filename": get_next_filename()})


@app.route("/api/build", methods=["POST"])
def api_build():
    data      = request.get_json(force=True)
    hymn_list = data.get("hymn_list", [])
    onegai    = data.get("onegai", "nashi")
    out_name  = data.get("filename", get_next_filename())
    if not out_name.lower().endswith(".pptx"):
        out_name += ".pptx"

    if not hymn_list:
        return jsonify({"ok": False, "error": "曲が指定されていません。"}), 400

    santabi_map = {}
    for item in get_santabi_items():
        fname = (item["prefix"] + "_" + item["title"] if item["prefix"] else item["title"]) + ".pptx"
        santabi_map[item["title"]] = os.path.join(SANTABI_DIR, fname)

    onegai_fname = "お願い証あり01.pptx" if onegai == "ari" else "お願い証なし01.pptx"
    onegai_path  = os.path.join(DATA_DIR, onegai_fname)

    prs = Presentation()
    prs.slide_width  = Inches(10)
    prs.slide_height = Inches(7.5)

    logs = []

    # お願いスライド（先頭）
    if os.path.exists(onegai_path):
        op = Presentation(onegai_path)
        if op.slides:
            copy_slide(op.slides[0], prs, op)
            logs.append(f"✔ お願いスライド追加: {onegai_fname}")
    else:
        logs.append(f"【警告】{onegai_fname} が data/ フォルダにありません。スキップします。")

    add_black_slide(prs)
    logs.append("▮ 黒スライドを挿入")

    for item in hymn_list:
        col_type = item.get("type")
        ident    = item.get("id")

        if col_type == "seika":
            n   = int(ident)
            src = get_seika_file(n)
            label = f"聖歌{n}番"
            if not src or not os.path.exists(src):
                logs.append(f"【エラー】{label} の元データが見つかりません。スキップ。")
                continue
            sp    = Presentation(src)
            found = [i for i, sl in enumerate(sp.slides)
                     if extract_hymn_number(slide_text(sl)) == n]
            logs.append(f"  {label}: {len(found)}枚")
            for i in found:
                sl = copy_slide(sp.slides[i], prs)
                set_bg(sl, NAVY)

        elif col_type == "kyusan":
            n   = int(ident)
            src = get_kyusan_file(n)
            label = f"旧賛美{n}番"
            if not src or not os.path.exists(src):
                logs.append(f"【エラー】{label} の元データが見つかりません。スキップ。")
                continue
            sp    = Presentation(src)
            found = [i for i, sl in enumerate(sp.slides)
                     if extract_kyusan_number(slide_text(sl)) == n]
            logs.append(f"  {label}: {len(found)}枚")
            for i in found:
                sl = copy_slide(sp.slides[i], prs)
                set_bg(sl, NAVY)

        else:  # santabi
            title = str(ident)
            src   = santabi_map.get(title)
            label = f"賛美集外「{title}」"
            if not src or not os.path.exists(src):
                logs.append(f"【エラー】{label} が見つかりません。スキップ。")
                continue
            sp = Presentation(src)
            logs.append(f"  {label}: {len(sp.slides)}枚")
            for sl in sp.slides:
                new_sl = copy_slide(sl, prs)
                set_bg(new_sl, BLUE)

        add_black_slide(prs)
        logs.append("▮ 黒スライドを挿入")

    out_path = os.path.join(OUTPUT_DIR, out_name)
    prs.save(out_path)
    total = len(prs.slides)
    logs.append(f"✔ 完了！ 合計 {total} 枚 → {out_name}")

    return jsonify({"ok": True, "filename": out_name, "logs": logs})


@app.route("/api/download/<path:filename>")
def api_download(filename):
    filename = os.path.basename(filename)
    path = os.path.join(OUTPUT_DIR, filename)
    if not os.path.exists(path):
        return "Not found", 404
    return send_file(
        path, as_attachment=True, download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5002, debug=False)
